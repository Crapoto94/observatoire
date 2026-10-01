# -*- coding: utf-8 -*-
"""Construit server/seed/indicators.json à partir du classeur Excel et de la carte mentale PDF.

- Excel : une cellule (contexte/suivi/évaluation/prospective) = 1 ou plusieurs indicateurs (1 par ligne de texte)
- PDF   : fournit la priorité (couleur du texte P1..P4) par rapprochement de libellés
Usage : python tools/build_seed.py
"""
import json, re, unicodedata, difflib, os
import openpyxl, pymupdf

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, 'indicateurs_rempli_Revu ON 22_09_2026.xlsx')
PDF = os.path.join(ROOT, '20261002_Carte_mentale_Observatoire_v1.pdf')
OUT = os.path.join(ROOT, 'server', 'seed', 'indicators.json')

THEMES = {  # onglet -> (clé, libellé)
    'demographie': ('demographie', 'Démographie'),
    'emploi economie': ('emploi', 'Emploi & économie'),
    'cohesion sociale santé': ('cohesion', 'Cohésion sociale & santé'),
    'Mobilite': ('mobilite', 'Mobilité'),
    'Logement urbanisme': ('logement', 'Logement & urbanisme'),
    'Env et STE': ('environnement', 'Environnement & STE'),
}
# libellé de groupe Excel (colonne A) -> (clé groupe, libellé carte)
GROUPS = {
    'Démographie et population': ('demographie', 'Démographie'),
    'Emploi, revenu, économie': ('emploi-revenus', 'Emploi et revenus'),
    'Commerces et développement économique': ('commerces', 'Commerces et dével. éco.'),
    'Conditions de vie cohésion sociale': ('conditions-vie', 'Conditions de vie'),
    'Vie associative': ('vie-associative', 'Vie associative'),
    'Santé': ('sante', 'Santé'),
    'Mobilités': ('mobilite', 'Mobilité'),
    'Logement et urbanisme': ('logement', 'Logement & urbanisme'),
    'Environnement, STE': ('environnement', 'Environnement & STE'),
}
LEVELS = ['contexte', 'suivi', 'evaluation', 'prospective']


def norm(s):
    s = unicodedata.normalize('NFKD', s.casefold())
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()


def clean(s):
    return re.sub(r'\s+', ' ', s.replace('’', "'").replace(' ', ' ')).strip()


# ---------------- Excel ----------------
wb = openpyxl.load_workbook(XLSX)
inds = []
for ws in wb:
    theme, theme_label = THEMES[ws.title]
    header = {c.column: str(c.value or '') for c in ws[1]}
    per_col = next((c for c, v in header.items() if 'PERIODICITE' in v.upper() or 'DATE/ANNEE' in v.upper()), 9)
    group = None
    order = 0
    for r in range(2, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a and clean(str(a)) in GROUPS:
            group = clean(str(a))
        if not group:
            continue
        cells = [ws.cell(r, 3 + i).value for i in range(4)]
        if not any(cells):
            continue
        lines = [[clean(x) for x in str(v).split('\n') if clean(x)] if v else [] for v in cells]
        n = max(len(l) for l in lines)
        src = clean(str(ws.cell(r, 7).value or ''))
        lien = clean(str(ws.cell(r, 8).value or ''))
        per = clean(str(ws.cell(r, per_col).value or ''))
        gkey, glabel = GROUPS[group]
        for k in range(n):
            for li, lvl in enumerate(LEVELS):
                if k < len(lines[li]):
                    inds.append(dict(
                        theme=theme, theme_label=theme_label, groupe=gkey, groupe_label=glabel,
                        excel_sheet=ws.title, excel_row=r, sous_ligne=k, ordre=order, niveau=lvl,
                        libelle=lines[li][k], source=src, lien_origine=lien, periodicite=per))
            order += 1

# ---------------- PDF : priorités ----------------
PRIO = {0xcc0000: 1, 0xff6666: 2, 0x66b2ff: 3, 0x0066cc: 4}
FILL = {(0.84, 0.91, 0.83): 'contexte', (1.0, 0.9, 0.8): 'suivi', (0.88, 0.84, 0.91): 'evaluation', (1.0, 0.95, 0.8): 'prospective'}
# zones de la carte (coordonnées PDF) -> clé de groupe
ZONES = [
    ('environnement', (100, 240, 1000, 960)),
    ('demographie', (1060, 20, 1830, 780)),
    ('emploi-revenus', (2380, 60, 3360, 570)),
    ('commerces', (2380, 575, 3360, 1075)),
    ('logement', (90, 1170, 1080, 2020)),
    ('mobilite', (1080, 1490, 1830, 2030)),
    ('conditions-vie', (2380, 1060, 3360, 1560)),
    ('vie-associative', (2380, 1561, 3360, 1800)),
    ('sante', (2380, 1801, 3360, 2110)),
]
doc = pymupdf.open(PDF)
page = doc[0]
rects = []
for d in page.get_drawings():
    f = d.get('fill')
    if f is None:
        continue
    key = tuple(round(c, 2) for c in f)
    if key in FILL and d['rect'].width > 100 and d['rect'].height > 30:
        rects.append((pymupdf.Rect(d['rect']), FILL[key]))
spans = [s for b in page.get_text('dict')['blocks'] for l in b.get('lines', []) for s in l['spans'] if s['text'].strip()]
cells = []
for rect, lvl in rects:
    inside = [s for s in spans if rect.contains(pymupdf.Point((s['bbox'][0] + s['bbox'][2]) / 2, (s['bbox'][1] + s['bbox'][3]) / 2))]
    if not inside:
        continue
    inside.sort(key=lambda s: (round(s['bbox'][1]), s['bbox'][0]))
    text = clean(' '.join(s['text'] for s in inside))
    prio = PRIO.get(inside[0]['color'])
    cx, cy = (rect.x0 + rect.x1) / 2, (rect.y0 + rect.y1) / 2
    zone = next((z for z, (x0, y0, x1, y1) in ZONES if x0 <= cx <= x1 and y0 <= cy <= y1), None)
    cells.append(dict(text=text, niveau=lvl, priorite=prio, zone=zone, norm=norm(text)))


def score(a, b):
    r = difflib.SequenceMatcher(None, a, b).ratio()
    ta, tb = set(a.split()), set(b.split())
    j = len(ta & tb) / max(1, len(ta | tb))
    return max(r, (r + j) / 2 + 0.05 if j > 0.6 else r)


pairs = []
for ci, c in enumerate(cells):
    for ii, i in enumerate(inds):
        if c['niveau'] == i['niveau'] and c['zone'] == i['groupe']:
            pairs.append((score(c['norm'], norm(i['libelle'])), ci, ii))
pairs.sort(reverse=True)
used_c, used_i = set(), set()
for s, ci, ii in pairs:
    if s < 0.55 or ci in used_c or ii in used_i:
        continue
    used_c.add(ci)
    used_i.add(ii)
    inds[ii]['priorite'] = cells[ci]['priorite']
    inds[ii]['libelle_carte'] = cells[ci]['text']

LEGEND = {'suivi', 'evaluation', 'prospectives', 'contexte'}


def second_pass(same_zone, same_level, threshold):
    cand = []
    for ci, c in enumerate(cells):
        if ci in used_c or c['norm'] in LEGEND or c['norm'].startswith('priorite'):
            continue
        for ii, i in enumerate(inds):
            if ii in used_i:
                continue
            if same_zone and c['zone'] != i['groupe']:
                continue
            if same_level and c['niveau'] != i['niveau']:
                continue
            cand.append((score(c['norm'], norm(i['libelle'])), ci, ii))
    cand.sort(reverse=True)
    for s, ci, ii in cand:
        if s < threshold or ci in used_c or ii in used_i:
            continue
        used_c.add(ci)
        used_i.add(ii)
        inds[ii]['priorite'] = cells[ci]['priorite']
        inds[ii]['libelle_carte'] = cells[ci]['text']
        if cells[ci]['niveau'] != inds[ii]['niveau']:
            inds[ii]['niveau_carte'] = cells[ci]['niveau']


second_pass(True, True, 0.35)    # même groupe et niveau, seuil bas (ex. "vacance" / "taux de vacance des logements")
second_pass(True, False, 0.5)    # même groupe, niveau différent (incohérence de colonne)
second_pass(False, True, 0.5)    # zone mal détectée

for n, i in enumerate(inds, 1):
    i['id'] = n
    i.setdefault('priorite', None)
    i.setdefault('libelle_carte', None)

json.dump(inds, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('indicateurs Excel :', len(inds), '| cellules carte :', len(cells))
print('apparies :', len(used_i), '| carte non appariees :', len(cells) - len(used_c), '| Excel sans priorite :', len(inds) - len(used_i))
print('--- cellules carte non appariees')
for ci, c in enumerate(cells):
    if ci not in used_c:
        print('  ', c['zone'], c['niveau'], c['priorite'], '|', c['text'])
print('--- indicateurs Excel non apparies')
for ii, i in enumerate(inds):
    if ii not in used_i:
        print('  ', i['groupe'], i['niveau'], '|', i['libelle'])
